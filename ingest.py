"""Ingest every file in the `information/` folder into data.csv.

For each file the script extracts text based on its type:
  - PDFs                -> pdfplumber
  - Excel spreadsheets  -> openpyxl
  - Word documents      -> python-docx
  - Plain text / minutes-> read directly
  - Images              -> OCR via EasyOCR (pure pip install, no native binary)

Each file becomes one row in data.csv with:
  - a one-sentence `summary`
  - AI-extracted `keywords` (comma separated)

Keyword extraction uses an LLM when an API key is configured
(GROQ_API_KEY, or GOOGLE_API_KEY / GEMINI_API_KEY), and automatically falls
back to a dependency-free local extractor otherwise, so the pipeline always
produces keywords and never hard-fails on a missing key or network.

Optimizations:
  - Incremental cache: files whose path+size+mtime are unchanged since the
    last run are reused from the existing data.csv (no re-OCR, no re-LLM).
  - Heavy libraries (EasyOCR/torch, LLM clients) are imported lazily, only
    when actually needed.
  - Atomic, lock-safe CSV write (temp file + replace) with a clear message
    if the output file is locked by another program (e.g. Excel).

Usage:
  python ingest.py              # incremental, AI keywords if a key is set
  python ingest.py --rebuild    # ignore cache, reprocess everything
  python ingest.py --no-ai      # force local keyword extraction only
"""

from __future__ import annotations

import csv
import datetime
import os
import re
import sys
import tempfile

INFORMATION_DIR = "information"
CSV_FILE = "data.csv"

MAX_SUMMARY = 240          # chars for the one-sentence summary
MAX_KEYWORDS = 8           # keywords stored per file
LLM_INPUT_CHARS = 6000     # how much text to send to the LLM per file

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif"}
EXCEL_EXTS = {".xlsx", ".xlsm", ".xls"}
TEXT_EXTS = {".txt", ".md", ".csv"}

FIELDNAMES = [
    "doc_id",
    "file_name",
    "relative_path",
    "doc_type",
    "size_bytes",
    "modified",
    "status",
    "summary",
    "keywords",
]

# Lazily-built singletons (expensive to construct, reused across files).
_EASYOCR_READER = None
_EASYOCR_ERROR = None
_LLM = None  # tuple (provider, client) or False once resolved


# --- Classification & text reduction --------------------------------------


def classify(ext: str) -> str:
    """Map a file extension to a human-friendly document category."""
    ext = ext.lower()
    if ext == ".pdf":
        return "PDF"
    if ext in EXCEL_EXTS:
        return "Spreadsheet"
    if ext == ".docx":
        return "Document"
    if ext in IMAGE_EXTS:
        return "Image"
    if ext in TEXT_EXTS:
        return "Text"
    return "Other"


def summarize_text(text: str) -> str:
    """Reduce extracted text to a single representative sentence."""
    if not text:
        return ""
    text = " ".join(text.split())
    for sentence in re.split(r"(?<=[.!?])\s+", text):
        candidate = sentence.strip()
        if len(candidate) >= 15 and re.search(r"[A-Za-z]", candidate):
            summary = candidate
            break
    else:
        summary = text
    if len(summary) > MAX_SUMMARY:
        summary = summary[: MAX_SUMMARY - 3].rstrip() + "..."
    return summary


# --- Keyword extraction (AI with local fallback) ---------------------------

_STOPWORDS = set(
    """
    a an the and or but if then else for to of in on at by with from as is are was were
    be been being this that these those it its it's into over under about above below
    not no nor so than too very can will just don should now i you he she we they them
    me my your our their his her out up down off again further once here there all any
    both each few more most other some such only own same up out which who whom what
    when where why how do does did doing have has had having would could may might must
    page version date name number total based using use used shall per via etc
    """.split()
)


def _resolve_llm():
    """Return (provider, client) if an API key is configured, else None."""
    global _LLM
    if _LLM is not None:
        return _LLM or None

    try:
        from dotenv import load_dotenv

        load_dotenv()
    except Exception:  # noqa: BLE001
        pass

    groq_key = os.environ.get("GROQ_API_KEY")
    google_key = os.environ.get("GOOGLE_API_KEY") or os.environ.get("GEMINI_API_KEY")

    if groq_key:
        try:
            from groq import Groq

            _LLM = ("groq", Groq(api_key=groq_key))
            return _LLM
        except Exception:  # noqa: BLE001
            pass

    if google_key:
        try:
            from google import genai

            _LLM = ("gemini", genai.Client(api_key=google_key))
            return _LLM
        except Exception:  # noqa: BLE001
            pass

    _LLM = False
    return None


def _parse_keywords(raw: str) -> list[str]:
    """Normalize an LLM/text response into a clean keyword list."""
    parts = re.split(r"[,\n;]+", raw)
    seen, out = set(), []
    for p in parts:
        kw = re.sub(r"^[\s\-\*\d\.\)]+", "", p).strip().strip('"').strip("'")
        kw = kw.lower()
        if 2 <= len(kw) <= 40 and kw not in seen:
            seen.add(kw)
            out.append(kw)
        if len(out) >= MAX_KEYWORDS:
            break
    return out


def _keywords_via_llm(text: str) -> list[str] | None:
    """Ask the configured LLM for keywords; return None if unavailable/failed."""
    llm = _resolve_llm()
    if not llm:
        return None
    provider, client = llm
    snippet = text[:LLM_INPUT_CHARS]
    prompt = (
        "Extract the most important keywords and key phrases from the document "
        f"text below. Return at most {MAX_KEYWORDS} items as a comma-separated "
        "list, lowercase, no numbering, no extra commentary.\n\n"
        f"TEXT:\n{snippet}"
    )
    try:
        if provider == "groq":
            resp = client.chat.completions.create(
                model="llama-3.1-8b-instant",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.2,
                max_tokens=120,
            )
            raw = resp.choices[0].message.content or ""
        else:  # gemini
            resp = client.models.generate_content(
                model="gemini-1.5-flash",
                contents=prompt,
            )
            raw = getattr(resp, "text", "") or ""
    except Exception:  # noqa: BLE001
        return None
    kws = _parse_keywords(raw)
    return kws or None


def _keywords_local(text: str) -> list[str]:
    """Dependency-free fallback: frequency-ranked content words/bigrams."""
    tokens = re.findall(r"[A-Za-z][A-Za-z\-']+", text.lower())
    tokens = [t for t in tokens if len(t) >= 3 and t not in _STOPWORDS]
    if not tokens:
        return []

    freq: dict[str, int] = {}
    for t in tokens:
        freq[t] = freq.get(t, 0) + 1

    for a, b in zip(tokens, tokens[1:]):
        bigram = f"{a} {b}"
        freq[bigram] = freq.get(bigram, 0) + 2  # weight phrases higher

    ranked = sorted(freq.items(), key=lambda kv: (-kv[1], kv[0]))
    out: list[str] = []
    for term, _count in ranked:
        if any(term != kw and term in kw.split() for kw in out):
            continue
        out.append(term)
        if len(out) >= MAX_KEYWORDS:
            break
    return out


def extract_keywords(text: str, use_ai: bool) -> tuple[str, str]:
    """Return (keywords_csv, method) using AI when enabled, else local."""
    if not text or not text.strip():
        return "", "none"
    if use_ai:
        kws = _keywords_via_llm(text)
        if kws is not None:
            return ", ".join(kws), "ai"
    return ", ".join(_keywords_local(text)), "local"


# --- Extractors ------------------------------------------------------------
# Each returns a tuple: (extracted_text, status_message)


def extract_pdf(path: str):
    try:
        import pdfplumber
    except ImportError:
        return "", "skipped: pdfplumber not installed (pip install pdfplumber)"
    try:
        parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        return "\n".join(parts), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract_excel(path: str):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "", "skipped: openpyxl not installed (pip install openpyxl)"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract_docx(path: str):
    try:
        import docx
    except ImportError:
        return "", "skipped: python-docx not installed (pip install python-docx)"
    try:
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract_text(path: str):
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read(), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def _get_easyocr_reader():
    """Lazily build and cache a single EasyOCR reader (English)."""
    global _EASYOCR_READER, _EASYOCR_ERROR
    if _EASYOCR_READER is not None or _EASYOCR_ERROR is not None:
        return _EASYOCR_READER
    try:
        import easyocr
    except ImportError:
        _EASYOCR_ERROR = "skipped: easyocr not installed (pip install easyocr)"
        return None
    try:
        _EASYOCR_READER = easyocr.Reader(["en"], gpu=False, verbose=False)
    except Exception as e:  # noqa: BLE001
        _EASYOCR_ERROR = f"error building EasyOCR reader: {e}"
        return None
    return _EASYOCR_READER


def extract_image(path: str):
    reader = _get_easyocr_reader()
    if reader is None:
        return "", _EASYOCR_ERROR or "skipped: OCR unavailable"
    try:
        results = reader.readtext(path, detail=0, paragraph=True)
        return " ".join(results), "ok"
    except Exception as e:  # noqa: BLE001
        return "", f"error: {e}"


def extract(path: str, ext: str):
    ext = ext.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in EXCEL_EXTS:
        return extract_excel(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in IMAGE_EXTS:
        return extract_image(path)
    if ext in TEXT_EXTS:
        return extract_text(path)
    return "", "skipped: unsupported file type"


# --- Incremental cache ------------------------------------------------------


def _load_cache(out: str) -> dict:
    """Index existing rows by relative_path for incremental reuse."""
    cache: dict[str, dict] = {}
    if not os.path.exists(out):
        return cache
    try:
        with open(out, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if row.get("relative_path"):
                    cache[row["relative_path"]] = row
    except Exception:  # noqa: BLE001 - a bad/old cache is simply ignored
        return {}
    return cache


def _cache_hit(cached: dict, stat, status_ok_only: bool = True) -> bool:
    """A cached row is reusable if size+mtime match and it was successful."""
    if not cached:
        return False
    try:
        same_size = str(stat.st_size) == str(cached.get("size_bytes"))
        same_time = cached.get("modified") == datetime.datetime.fromtimestamp(
            stat.st_mtime
        ).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:  # noqa: BLE001
        return False
    if status_ok_only and cached.get("status") != "ok":
        return False
    return same_size and same_time


# --- Pipeline --------------------------------------------------------------


def iter_files(root: str):
    for dirpath, _dirs, files in os.walk(root):
        for name in sorted(files):
            if name.lower() == "readme.txt":
                continue
            yield os.path.join(dirpath, name)


def _atomic_write(out: str, rows: list[dict]) -> None:
    """Write CSV via a temp file then replace, so a locked target fails cleanly."""
    out_dir = os.path.dirname(os.path.abspath(out)) or "."
    fd, tmp = tempfile.mkstemp(prefix=".ingest-", suffix=".csv", dir=out_dir)
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
            writer.writeheader()
            writer.writerows(rows)
        os.replace(tmp, out)  # atomic on the same volume
    except PermissionError as e:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise PermissionError(
            f"Cannot write '{out}' - it may be open in another program "
            f"(e.g. Excel). Close it and re-run. ({e})"
        ) from e
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise


def ingest(root=INFORMATION_DIR, out=CSV_FILE, use_ai=True, rebuild=False):
    os.makedirs(root, exist_ok=True)

    cache = {} if rebuild else _load_cache(out)
    rows = []
    reused = 0

    for doc_id, path in enumerate(iter_files(root), start=1):
        ext = os.path.splitext(path)[1]
        rel = os.path.relpath(path, root).replace("\\", "/")
        stat = os.stat(path)

        cached = cache.get(rel)
        if _cache_hit(cached, stat):
            row = dict(cached)
            row["doc_id"] = doc_id  # keep ids contiguous for the new ordering
            rows.append(row)
            reused += 1
            continue

        text, status = extract(path, ext)
        keywords, _method = extract_keywords(text, use_ai=use_ai)
        rows.append(
            {
                "doc_id": doc_id,
                "file_name": os.path.basename(path),
                "relative_path": rel,
                "doc_type": classify(ext),
                "size_bytes": stat.st_size,
                "modified": datetime.datetime.fromtimestamp(
                    stat.st_mtime
                ).strftime("%Y-%m-%d %H:%M:%S"),
                "status": status,
                "summary": summarize_text(text),
                "keywords": keywords,
            }
        )

    _atomic_write(out, rows)
    return rows, reused


def main(argv=None):
    argv = argv if argv is not None else sys.argv[1:]
    use_ai = "--no-ai" not in argv
    rebuild = "--rebuild" in argv

    rows, reused = ingest(use_ai=use_ai, rebuild=rebuild)

    provider = _resolve_llm()
    kw_mode = f"AI ({provider[0]})" if (use_ai and provider) else "local"
    ok = sum(1 for r in rows if r["status"] == "ok")

    print(f"Ingested {len(rows)} file(s) into {CSV_FILE}")
    print(f"  extracted ok: {ok}")
    print(f"  reused from cache: {reused}")
    print(f"  keyword mode: {kw_mode}")
    for r in rows:
        if r["status"] != "ok":
            print(f"    - {r['file_name']}: {r['status']}")


if __name__ == "__main__":
    main()
